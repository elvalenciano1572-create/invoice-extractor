"""Write invoices to Excel or CSV."""

from __future__ import annotations

import csv
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from .models import Invoice

COLUMNS = (
    ("File", "source"),
    ("Vendor", "vendor"),
    ("Invoice no.", "number"),
    ("Date", "issue_date"),
    ("Subtotal", "subtotal"),
    ("Tax %", "tax_rate"),
    ("Tax", "tax"),
    ("Total", "total"),
    ("Currency", "currency"),
    ("Status", None),
    ("Issues", None),
)
MONEY = {"Subtotal", "Tax", "Total"}
_HEADER = Font(bold=True, color="FFFFFF")
_HEADER_FILL = PatternFill("solid", fgColor="305496")
_CHECK_FILL = PatternFill("solid", fgColor="FCE4D6")


def _row(inv: Invoice) -> list:
    row = []
    for header, attr in COLUMNS:
        if header == "Status":
            row.append("OK" if inv.ok else "CHECK")
        elif header == "Issues":
            row.append("; ".join(inv.issues))
        else:
            value = getattr(inv, attr)
            row.append(float(value) if isinstance(value, Decimal) else value)
    return row


def to_xlsx(invoices: list[Invoice], path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Invoices"
    ws.append([h for h, _ in COLUMNS])
    for cell in ws[1]:
        cell.font, cell.fill = _HEADER, _HEADER_FILL

    for inv in invoices:
        ws.append(_row(inv))
        if not inv.ok:
            for cell in ws[ws.max_row]:
                cell.fill = _CHECK_FILL

    last = ws.max_row
    headers = [h for h, _ in COLUMNS]
    cur_col = get_column_letter(headers.index("Currency") + 1)
    if last > 1:
        ws.append([])
        # One total per currency: adding euros and dollars together means nothing.
        currencies = sorted({inv.currency for inv in invoices if inv.currency})
        for code in currencies:
            ws.append([f"TOTAL {code}"] + [None] * (len(COLUMNS) - 1))
            row = ws.max_row
            for i, header in enumerate(headers, start=1):
                col = get_column_letter(i)
                if header in MONEY:
                    ws[f"{col}{row}"] = (
                        f'=SUMIF({cur_col}2:{cur_col}{last},"{code}",{col}2:{col}{last})'
                    )
            ws[f"{cur_col}{row}"] = code
            for cell in ws[row]:
                cell.font = Font(bold=True)

        for i, header in enumerate(headers, start=1):
            col = get_column_letter(i)
            if header in MONEY or header == "Tax %":
                for r in range(2, ws.max_row + 1):
                    ws[f"{col}{r}"].number_format = "#,##0.00"
            if header == "Date":
                for r in range(2, last + 1):
                    ws[f"{col}{r}"].number_format = "yyyy-mm-dd"

    widths = {"File": 32, "Vendor": 28, "Invoice no.": 16, "Date": 12, "Issues": 60}
    for i, (header, _) in enumerate(COLUMNS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = widths.get(header, 12)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}{last}"
    wb.save(path)


def to_csv(invoices: list[Invoice], path: Path) -> None:
    with path.open("w", newline="", encoding="utf-8") as fh:
        writer = csv.writer(fh)
        writer.writerow([h for h, _ in COLUMNS])
        for inv in invoices:
            writer.writerow(_row(inv))
