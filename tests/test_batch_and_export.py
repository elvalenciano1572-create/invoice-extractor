import csv
from decimal import Decimal

from openpyxl import load_workbook

from invoice_extractor.batch import process_folder
from invoice_extractor.cli import main
from invoice_extractor.export import COLUMNS, to_csv, to_xlsx


def by_file(invoices):
    return {inv.source: inv for inv in invoices}


def test_folder_end_to_end(samples_dir):
    invoices = by_file(process_folder(samples_dir))
    assert len(invoices) == 6

    assert invoices["northwind_office.pdf"].ok
    assert invoices["northwind_office.pdf"].total == Decimal("2690.80")

    turia = invoices["imprenta_turia.pdf"]
    assert turia.number == "2026/317"
    assert any("!= total" in i for i in turia.issues)

    for name in ("suministros_levante.pdf", "suministros_levante_reenvio.pdf"):
        assert any("duplicate" in i for i in invoices[name].issues)

    assert "missing issue_date" in invoices["cafeteria_ruzafa.pdf"].issues
    assert any("no text layer" in i for i in invoices["escaneada.pdf"].issues)


def test_broken_file_does_not_stop_batch(tmp_path):
    (tmp_path / "broken.pdf").write_bytes(b"not a pdf")
    invoices = process_folder(tmp_path)
    assert len(invoices) == 1
    assert any("could not read" in i for i in invoices[0].issues)


def test_xlsx_export(samples_dir, tmp_path):
    out = tmp_path / "out.xlsx"
    to_xlsx(process_folder(samples_dir), out)
    ws = load_workbook(out)["Invoices"]
    assert [c.value for c in ws[1]] == [h for h, _ in COLUMNS]
    assert ws.max_row == 1 + 6 + 1 + 2  # header, invoices, blank, one total per currency
    total_col = [h for h, _ in COLUMNS].index("Total") + 1
    totals = {
        ws.cell(r, 1).value: ws.cell(r, total_col).value for r in (ws.max_row - 1, ws.max_row)
    }
    assert set(totals) == {"TOTAL EUR", "TOTAL USD"}, "currencies must never be summed together"
    assert '"EUR"' in totals["TOTAL EUR"] and totals["TOTAL EUR"].startswith("=SUMIF(")
    statuses = {ws.cell(r, 1).value: ws.cell(r, 10).value for r in range(2, 8)}
    assert statuses["northwind_office.pdf"] == "OK"
    assert statuses["imprenta_turia.pdf"] == "CHECK"


def test_csv_export(samples_dir, tmp_path):
    out = tmp_path / "out.csv"
    to_csv(process_folder(samples_dir), out)
    rows = list(csv.reader(out.open(encoding="utf-8")))
    assert rows[0][0] == "File"
    assert len(rows) == 7


def test_cli_exit_codes(samples_dir, tmp_path, capsys):
    out = tmp_path / "r.xlsx"
    assert main([str(samples_dir), "-o", str(out)]) == 1  # some invoices need checking
    assert out.exists()
    assert "6 invoices" in capsys.readouterr().out

    assert main([str(tmp_path / "missing")]) == 2
    empty = tmp_path / "empty"
    empty.mkdir()
    assert main([str(empty)]) == 2


def test_cli_all_ok(tmp_path, samples_dir):
    clean = tmp_path / "clean"
    clean.mkdir()
    (clean / "n.pdf").write_bytes((samples_dir / "northwind_office.pdf").read_bytes())
    assert main([str(clean), "-o", str(tmp_path / "x.csv")]) == 0
